import csv
import json

import pytest
from openpyxl import load_workbook

from jobscraper.models import COLUMNS, Job
from jobscraper.output import infer_format, write_jobs

JOBS = [
    Job(title="=HYPERLINK(\"http://evil\")", url="https://jobs.example.com/1", company="Acme\x01",
        certifications=["CISSP", "OSCP"], requirements=["Degree", "3+ years"], description="Full text"),
    Job(title="SOC Analyst", url="https://jobs.example.com/2", company="  =cmd()"),
]


def test_infer_format():
    assert infer_format("out.xlsx") == "xlsx"
    assert infer_format("out.TSV") == "tsv"
    assert infer_format("out") == "csv"
    assert infer_format("out.csv", "json") == "json"
    with pytest.raises(ValueError):
        infer_format("x", "pdf")


def test_csv_output(tmp_path):
    path = tmp_path / "jobs.csv"
    write_jobs(JOBS, str(path))
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig")))
    assert list(rows[0]) == COLUMNS
    assert rows[0]["title"].startswith("'=")  # formula injection neutralised
    assert rows[0]["company"] == "Acme"
    assert rows[0]["certifications"] == "CISSP; OSCP"
    assert rows[0]["requirements"] == "Degree | 3+ years"
    assert rows[1]["url"] == "https://jobs.example.com/2"
    assert rows[1]["company"] == "'  =cmd()"


def test_tsv_and_description(tmp_path):
    path = tmp_path / "jobs.tsv"
    write_jobs(JOBS, str(path), include_description=True)
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig"), delimiter="\t"))
    assert rows[0]["description"] == "Full text"


def test_json_output(tmp_path):
    path = tmp_path / "jobs.json"
    write_jobs(JOBS, str(path))
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data[0]["certifications"] == ["CISSP", "OSCP"]
    assert "description" not in data[0]


def test_xlsx_output(tmp_path):
    path = tmp_path / "jobs.xlsx"
    write_jobs(JOBS, str(path))
    sheet = load_workbook(path).active
    assert [c.value for c in sheet[1]] == COLUMNS
    title = sheet.cell(row=2, column=1)
    assert title.data_type == "s" and title.value.startswith("'=")
    url = sheet.cell(row=2, column=COLUMNS.index("url") + 1)
    assert url.hyperlink.target == "https://jobs.example.com/1"
    with pytest.raises(ValueError):
        write_jobs(JOBS, "-", "xlsx")
