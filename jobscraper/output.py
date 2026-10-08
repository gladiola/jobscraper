"""Write jobs to CSV, TSV, Excel (XLSX) or JSON."""

from __future__ import annotations

import csv
import json
import os
import re
import sys
from typing import Any, Iterable, List, Optional, Sequence

from .models import COLUMNS, Job

FORMATS = ("csv", "tsv", "xlsx", "json")
_EXTENSIONS = {".csv": "csv", ".tsv": "tsv", ".tab": "tsv", ".xlsx": "xlsx", ".json": "json"}
_LIST_SEPARATORS = {"requirements": " | "}
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")
_EXCEL_MAX_CELL = 32767
# Control characters that are invalid in XLSX (and useless in CSV).
_CONTROL_CHARS = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def infer_format(path: str, fmt: Optional[str] = None) -> str:
    """Return the output format, from ``fmt`` or from the file extension (default CSV)."""
    if fmt:
        fmt = fmt.lower()
        if fmt not in FORMATS:
            raise ValueError(f"Unsupported format {fmt!r}; choose from: {', '.join(FORMATS)}")
        return fmt
    return _EXTENSIONS.get(os.path.splitext(path)[1].lower(), "csv")


def _cell(column: str, value: Any) -> str:
    """Flatten a value to a spreadsheet-safe string."""
    if isinstance(value, (list, tuple)):
        value = _LIST_SEPARATORS.get(column, "; ").join(str(v) for v in value)
    text = _CONTROL_CHARS.sub("", "" if value is None else str(value))
    # Scraped text is untrusted: stop spreadsheet apps from evaluating it as a formula.
    if text.lstrip(" ").startswith(_FORMULA_PREFIXES):
        text = "'" + text
    return text


def _rows(jobs: Iterable[Job], columns: Sequence[str]) -> Iterable[List[str]]:
    for job in jobs:
        data = job.to_dict()
        yield [_cell(c, data.get(c)) for c in columns]


def write_jobs(jobs: Sequence[Job], path: str, fmt: Optional[str] = None, include_description: bool = False,
               append: bool = False) -> str:
    """Write ``jobs`` to ``path`` ("-" for stdout) and return the format used."""
    fmt = infer_format(path, fmt)
    columns = COLUMNS + (["description"] if include_description else [])

    if fmt == "xlsx":
        if path == "-":
            raise ValueError("XLSX output cannot be written to stdout; give a file name")
        _write_xlsx(jobs, path, columns, append)
        return fmt

    if fmt == "json":
        records = []
        for job in jobs:
            record = job.to_dict()
            if not include_description:
                record.pop("description", None)
            records.append(record)
        if append and path != "-" and os.path.isfile(path) and os.path.getsize(path):
            with open(path, encoding="utf-8") as existing:
                previous = json.load(existing)
            if not isinstance(previous, list):
                raise ValueError("Cannot append JSON output: existing file must contain a JSON array")
            records = previous + records
        stream = sys.stdout if path == "-" else open(path, "w", encoding="utf-8")
        try:
            json.dump(records, stream, indent=2, ensure_ascii=False)
            stream.write("\n")
        finally:
            if stream is not sys.stdout:
                stream.close()
        return fmt

    # utf-8-sig adds a byte-order mark so Excel detects UTF-8 in CSV/TSV files.
    has_output = append and path != "-" and os.path.isfile(path) and os.path.getsize(path) > 0
    encoding = "utf-8" if has_output else "utf-8-sig"
    mode = "a" if has_output else ("a" if append and path != "-" else "w")
    stream = sys.stdout if path == "-" else open(path, mode, newline="", encoding=encoding)
    try:
        writer = csv.writer(stream, delimiter="\t" if fmt == "tsv" else ",")
        if not has_output:
            writer.writerow(columns)
        writer.writerows(_rows(jobs, columns))
    finally:
        if stream is not sys.stdout:
            stream.close()
    return fmt


def _write_xlsx(jobs: Sequence[Job], path: str, columns: Sequence[str], append: bool = False) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    if append and os.path.isfile(path) and os.path.getsize(path):
        from openpyxl import load_workbook

        workbook = load_workbook(path)
        if "Jobs" not in workbook.sheetnames:
            raise ValueError("Cannot append XLSX output: existing workbook has no 'Jobs' sheet")
        sheet = workbook["Jobs"]
        if [cell.value for cell in sheet[1]] != list(columns):
            raise ValueError("Cannot append XLSX output: existing columns do not match")
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Jobs"
        sheet.append(list(columns))
        for cell in sheet[1]:
            cell.font = Font(bold=True)

    url_col = columns.index("url") + 1
    for row in _rows(jobs, columns):
        sheet.append([value[:_EXCEL_MAX_CELL] for value in row])
        link = sheet.cell(row=sheet.max_row, column=url_col)
        if str(link.value).startswith(("http://", "https://")):
            link.hyperlink = link.value
            link.style = "Hyperlink"

    widths = {"title": 40, "url": 50, "summary": 80, "requirements": 80, "description": 100}
    for index, column in enumerate(columns, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths.get(column, 20)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(path)
